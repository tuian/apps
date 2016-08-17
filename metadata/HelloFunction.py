import openpyxl
import warnings
import docx
import os.path as os

def checkFile(flag):
    #print "List of matching files:"
    warnings.simplefilter('ignore')

    parentDir = "C:\IFS\Automation_After_Corrections\\"
    input_excel_filename = "IFS Documents Analysis.xlsx"
    input_excel_filepath =  parentDir + input_excel_filename


    #read excel sheet for input
    try:

	wb = openpyxl.load_workbook(input_excel_filepath)#open input excel sheet
    except IOError:
	    print "ERROR: Could not open the input excel spreadsheet."
	    print "ERROR DETAILS: Input file '" + input_excel_filepath +"' is missing. Please check the file name and directory details."
    else:
	    #open sheet by name
	    sheet = wb.get_sheet_by_name('Sheet1')

    count = 0
    list_files_found = []
    list_files_not_found = []

    for row in range(2, (sheet.max_row+1)):
        if (sheet['C' + str(row)].value != None):
            if ((sheet['B' + str(row)].value).lower() == "Open".lower()):
                #print "\t File Name: " + sheet['C' + str(row)].value
                count=count+1

                # get name of ifs doc to read
                # add .docx extension
                docName = parentDir + sheet['C' + str(row)].value + '.docx'


                if(os.isfile(docName) == True):
                    #print "File exists"
                    list_files_found.append(docName)

                else:
                    #print "File does not exists"
                    list_files_not_found.append(docName)

                """
                try:
                    doc1 = docx.Document(docName)

                except:
                    print "ERROR in reading a document. Check if the file exists or the filename"
                    exit()
                """

    if(flag == 'FOUND'):
        return list_files_found

    if (flag == 'NOT_FOUND'):
        return list_files_not_found


def printListItems(flag):
    list = checkFile(flag)

    print "\nTotal Files - "+ flag +" : ", len(list)

    for i in range(0,len(list)):
        print list[i]



printListItems('FOUND')
printListItems('NOT_FOUND')

